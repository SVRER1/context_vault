import docx
from pathlib import Path
from contextvault.core.models import ParsedDocument, DocumentSection
from contextvault.core.exceptions import ParseError
from contextvault.parsers.base import BaseParser

class DOCXParser(BaseParser):
    """Parser for DOCX files using python-docx."""
    
    def supported_extensions(self) -> set[str]:
        return {'.docx'}
        
    def parse(self, file_path: Path, file_id: str) -> ParsedDocument:
        try:
            doc = docx.Document(str(file_path))
        except Exception as e:
            raise ParseError(f"Failed to open DOCX {file_path}: {e}")
            
        title = doc.core_properties.title
        if not title:
            title = file_path.name
            
        sections = []
        current_title = "Document Start"
        current_content = []
        current_heading_level = None
        current_char_start = 0
        current_char_end = 0
        current_paragraph_start = 1
        current_paragraph_end = 0
        full_content = []
        char_cursor = 0

        for paragraph_index, para in enumerate(doc.paragraphs, start=1):
            text = para.text.strip()
            if not text:
                continue
            if full_content:
                char_cursor += 2
            paragraph_char_start = char_cursor
            full_content.append(text)
            char_cursor += len(text)

            if para.style.name.startswith('Heading'):
                if current_content:
                    sections.append(DocumentSection(
                        heading=current_title,
                        text='\n'.join(current_content),
                        char_start=current_char_start,
                        char_end=current_char_end,
                        metadata={
                            "heading": current_title,
                            "paragraph_start": current_paragraph_start,
                            "paragraph_end": current_paragraph_end,
                        },
                    ))
                current_title = text
                current_content = [text]
                current_heading_level = int(para.style.name.removeprefix("Heading").strip() or 1)
                current_char_start = paragraph_char_start
                current_paragraph_start = paragraph_index
            else:
                current_content.append(text)
                if not current_content[:-1]:
                    current_char_start = paragraph_char_start
                    current_paragraph_start = paragraph_index
            current_char_end = char_cursor
            current_paragraph_end = paragraph_index

        if current_content:
            sections.append(DocumentSection(
                heading=current_title,
                text='\n'.join(current_content),
                char_start=current_char_start,
                char_end=current_char_end,
                level=current_heading_level,
                metadata={
                    "heading": current_title,
                    "paragraph_start": current_paragraph_start,
                    "paragraph_end": current_paragraph_end,
                },
            ))
            
        return ParsedDocument(
            file_id=file_id,
            path=str(file_path),
            title=title,
            text='\n\n'.join(full_content),
            sections=sections,
            metadata={"source_type": "docx"}
        )
