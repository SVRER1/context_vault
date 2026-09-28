import csv
import openpyxl
from openpyxl.utils import get_column_letter
from pathlib import Path
from contextvault.core.models import ParsedDocument, DocumentSection
from contextvault.core.exceptions import ParseError
from contextvault.parsers.base import BaseParser

class SpreadsheetParser(BaseParser):
    """Parser for CSV and XLSX files."""
    
    MAX_ROWS = 1000
    
    def supported_extensions(self) -> set[str]:
        return {'.csv', '.xlsx'}
        
    def parse(self, file_path: Path, file_id: str) -> ParsedDocument:
        ext = file_path.suffix.lower()
        
        if ext == '.csv':
            return self._parse_csv(file_path, file_id)
        elif ext == '.xlsx':
            return self._parse_xlsx(file_path, file_id)
        else:
            raise ParseError(f"Unsupported spreadsheet extension: {ext}")
            
    def _parse_csv(self, file_path: Path, file_id: str) -> ParsedDocument:
        content_lines = []
        row_count = 0
        max_columns = 0
        truncated = False
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                for i, row in enumerate(reader):
                    if i >= self.MAX_ROWS:
                        truncated = True
                        break
                    if any(row):
                        content_lines.append(" | ".join(row))
                    row_count = i + 1
                    max_columns = max(max_columns, len(row))
        except UnicodeDecodeError:
            try:
                with open(file_path, 'r', encoding='latin-1') as f:
                    reader = csv.reader(f)
                    for i, row in enumerate(reader):
                        if i >= self.MAX_ROWS:
                            truncated = True
                            break
                        if any(row):
                            content_lines.append(" | ".join(row))
                        row_count = i + 1
                        max_columns = max(max_columns, len(row))
            except Exception as e:
                raise ParseError(f"Failed to read CSV {file_path}: {e}")
        except Exception as e:
            raise ParseError(f"Failed to read CSV {file_path}: {e}")
            
        content = '\n'.join(content_lines)
        sections = []
        if content:
            sections.append(DocumentSection(
                heading="CSV Data",
                text=content,
                cell_range=f"A1:{get_column_letter(max(max_columns, 1))}{row_count}" if row_count else None,
                metadata={"sheet": "CSV Data"}
            ))
            
        return ParsedDocument(
            file_id=file_id,
            path=str(file_path),
            title=file_path.name,
            text=content,
            sections=sections,
            metadata={
                "source_type": "spreadsheet",
                "truncated": truncated,
                "rows_read": row_count,
            }
        )
        
    def _parse_xlsx(self, file_path: Path, file_id: str) -> ParsedDocument:
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except Exception as e:
            raise ParseError(f"Failed to open XLSX {file_path}: {e}")
            
        sections = []
        full_content = []
        truncated_sheets = []
        
        try:
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_lines = []
                row_count = 0
                
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i >= self.MAX_ROWS:
                        truncated_sheets.append(sheet_name)
                        break
                    row_str = [str(cell) if cell is not None else "" for cell in row]
                    if any(row_str):
                        sheet_lines.append(" | ".join(row_str))
                    row_count = i + 1
                        
                if sheet_lines:
                    sheet_content = '\n'.join(sheet_lines)
                    full_content.append(sheet_content)
                    sections.append(DocumentSection(
                        heading=sheet_name,
                        text=sheet_content,
                        sheet=sheet_name,
                        cell_range=f"A1:{get_column_letter(max(sheet.max_column, 1))}{row_count}" if row_count else None,
                        metadata={"sheet": sheet_name}
                    ))
        finally:
            wb.close()
            
        return ParsedDocument(
            file_id=file_id,
            path=str(file_path),
            title=file_path.name,
            text='\n\n'.join(full_content),
            sections=sections,
            metadata={
                "source_type": "spreadsheet",
                "truncated": bool(truncated_sheets),
                "truncated_sheets": truncated_sheets,
            }
        )
