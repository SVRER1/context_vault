import os

import docx
import fitz
import openpyxl
import pytest
from pptx import Presentation

from contextvault.core.config import AppConfig
from contextvault.core.vault import Vault
from contextvault.filesystem.operations import FileOperations
from contextvault.indexing import index_service as index_service_module
from contextvault.indexing.index_service import IndexService
from contextvault.storage.database import Database
from contextvault.services.vault_service import VaultService


def make_index(tmp_path):
    root = tmp_path / "vault"
    root.mkdir()
    config = AppConfig(app_data_dir=str(tmp_path / "app-data"))
    app_db = Database(config.app_db_path)
    app_db.initialize()
    vault_service = VaultService(config, app_db)
    info = vault_service.registry.register_vault(root, "test")
    vault = Vault(info)
    db = vault_service.get_vault_db(vault)
    return root, config, vault, db, app_db, IndexService(vault, db, config)


def close_index(db, app_db):
    db.close()
    app_db.close()


def test_second_reconcile_skips_hash_and_extraction(tmp_path, monkeypatch):
    root, _config, _vault, db, app_db, service = make_index(tmp_path)
    (root / "note.md").write_text("# Retrieval\nVector database notes.", encoding="utf-8")
    first = service.reconcile()
    original_hash = index_service_module.compute_sha256
    hash_calls = []
    extract_calls = []
    monkeypatch.setattr(index_service_module, "compute_sha256", lambda path: (hash_calls.append(path), original_hash(path))[1])
    original_extract = service.extractor.extract
    monkeypatch.setattr(service.extractor, "extract", lambda *args: (extract_calls.append(args[0]), original_extract(*args))[1])

    second = service.reconcile()

    assert first["created"] == 1
    assert second["unchanged"] == 1
    assert hash_calls == []
    assert extract_calls == []
    close_index(db, app_db)


def test_changed_bytes_update_existing_file_and_replace_fts(tmp_path):
    root, _config, _vault, db, app_db, service = make_index(tmp_path)
    path = root / "notes.txt"
    path.write_text("alpha beta", encoding="utf-8")
    service.reconcile()
    original = db.fetch_one("SELECT id FROM files WHERE filename='notes.txt'")["id"]
    old_mtime = path.stat().st_mtime
    path.write_text("gamma beta", encoding="utf-8")
    os.utime(path, (old_mtime + 5, old_mtime + 5))

    result = service.reconcile()

    assert result["updated"] == 1
    assert db.fetch_one("SELECT id FROM files WHERE filename='notes.txt'")["id"] == original
    assert db.fetch_one("SELECT id FROM chunks_fts WHERE text MATCH 'gamma'")
    assert db.fetch_one("SELECT id FROM chunks_fts WHERE text MATCH 'alpha'") is None
    close_index(db, app_db)


def test_unique_external_move_keeps_id_and_updates_passage_paths(tmp_path):
    root, _config, _vault, db, app_db, service = make_index(tmp_path)
    old = root / "before.md"
    old.write_text("A unique retrieval passage.", encoding="utf-8")
    service.reconcile()
    file_id = db.fetch_one("SELECT id FROM files WHERE filename='before.md'")["id"]
    new = root / "after.md"
    old.rename(new)

    result = service.reconcile()

    row = db.fetch_one("SELECT id, relative_path FROM files WHERE filename='after.md'")
    fts = db.fetch_one("SELECT relative_path, filename FROM chunks_fts WHERE file_id=?", (file_id,))
    assert result["moved"] == 1
    assert result["deleted"] == 0
    assert row["id"] == file_id
    assert row["relative_path"] == "after.md"
    assert fts["relative_path"] == "after.md"
    assert fts["filename"] == "after.md"
    close_index(db, app_db)


def test_operation_ledger_move_preserves_file_identity(tmp_path):
    root, _config, vault, db, app_db, service = make_index(tmp_path)
    old = root / "original.md"
    old.write_text("application move preserves this passage", encoding="utf-8")
    service.reconcile()
    original_id = db.fetch_one("SELECT id FROM files WHERE filename='original.md'")["id"]
    new = root / "renamed.md"
    FileOperations(db).move_file(old, new, vault, reason="test ledger move")

    result = service.reconcile()

    row = db.fetch_one("SELECT id FROM files WHERE filename='renamed.md'")
    assert result["moved"] == 1
    assert row["id"] == original_id
    assert db.fetch_one("SELECT relative_path FROM chunks_fts WHERE file_id=?", (original_id,))["relative_path"] == "renamed.md"
    close_index(db, app_db)


def test_delete_removes_file_chunks_and_fts(tmp_path):
    root, _config, _vault, db, app_db, service = make_index(tmp_path)
    path = root / "gone.txt"
    path.write_text("old searchable content", encoding="utf-8")
    service.reconcile()
    path.unlink()

    result = service.reconcile()

    assert result["deleted"] == 1
    assert db.fetch_one("SELECT id FROM files WHERE filename='gone.txt'") is None
    assert db.fetch_one("SELECT id FROM chunks WHERE text LIKE '%searchable%'") is None
    assert db.fetch_one("SELECT id FROM chunks_fts WHERE text MATCH 'searchable'") is None
    close_index(db, app_db)


def test_failed_traversal_does_not_prune_existing_index(tmp_path, monkeypatch):
    root, _config, _vault, db, app_db, service = make_index(tmp_path)
    (root / "keep.txt").write_text("keep this indexed", encoding="utf-8")
    service.reconcile()
    monkeypatch.setattr(service, "_discover", lambda _root: (_ for _ in ()).throw(PermissionError("fixture")))

    result = service.reconcile()

    assert result["status"] == "failed"
    assert db.fetch_one("SELECT id FROM files WHERE filename='keep.txt'") is not None
    assert db.fetch_one("SELECT status FROM index_runs ORDER BY started_at DESC LIMIT 1")["status"] == "failed"
    close_index(db, app_db)


def test_hash_failure_marks_partial_and_prevents_pruning(tmp_path, monkeypatch):
    root, _config, _vault, db, app_db, service = make_index(tmp_path)
    unreadable = root / "unreadable.txt"
    deleted = root / "not-scanned-after-error.txt"
    unreadable.write_text("old valid passage", encoding="utf-8")
    deleted.write_text("preserve until complete discovery", encoding="utf-8")
    service.reconcile()
    old_mtime = unreadable.stat().st_mtime
    unreadable.write_text("new content to hash", encoding="utf-8")
    os.utime(unreadable, (old_mtime + 5, old_mtime + 5))
    deleted.unlink()
    monkeypatch.setattr(
        index_service_module,
        "compute_sha256",
        lambda _path: (_ for _ in ()).throw(PermissionError("fixture denied")),
    )

    result = service.reconcile()

    assert result["status"] == "partial"
    assert result["failed"] == 1
    assert db.fetch_one("SELECT id FROM files WHERE filename=?", (deleted.name,)) is not None
    assert db.fetch_one("SELECT id FROM chunks_fts WHERE text MATCH 'valid'") is None
    assert db.fetch_one("SELECT extract_status FROM files WHERE filename=?", (unreadable.name,))["extract_status"] == "error"
    close_index(db, app_db)


def test_supported_formats_index_offline_without_model(tmp_path):
    root, _config, _vault, db, app_db, service = make_index(tmp_path)
    (root / "plain.txt").write_text("Plain text fixture passage.", encoding="utf-8")
    (root / "notes.md").write_text("# Markdown heading\nMarkdown structure fixture.", encoding="utf-8")
    (root / "data.json").write_text('{"topic": "jsonfixture", "value": 7}', encoding="utf-8")
    (root / "source.py").write_text("def codefixture():\n    return 'indexed'\n", encoding="utf-8")
    (root / "page.html").write_text("<html><body><h1>HtmlFixture</h1><p>Indexed content.</p></body></html>", encoding="utf-8")
    (root / "table.csv").write_text("topic,value\ncsvfixture,42\n", encoding="utf-8")

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((40, 60), "pdffixture indexed passage")
    pdf.save(root / "paper.pdf")
    pdf.close()

    document = docx.Document()
    document.add_heading("DocxFixture", level=1)
    document.add_paragraph("Indexed paragraph.")
    document.save(root / "word.docx")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PptxFixture"
    presentation.save(root / "slides.pptx")

    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "xlsxfixture"
    workbook.active["B1"] = "indexed"
    workbook.save(root / "sheet.xlsx")

    result = service.reconcile()

    assert result["status"] == "complete"
    assert result["failed"] == 0
    rows = db.fetch_all("SELECT filename, extract_status, word_count FROM files WHERE vault_id=?", (_vault.vault_id,))
    assert len(rows) == 10
    assert all(row["extract_status"] == "ok" and row["word_count"] > 0 for row in rows)
    for term in ("jsonfixture", "codefixture", "htmlfixture", "pdffixture", "docxfixture", "pptxfixture", "csvfixture", "xlsxfixture"):
        assert db.fetch_one("SELECT id FROM chunks_fts WHERE text MATCH ?", (term,)), term
    close_index(db, app_db)


def test_reconcile_ignores_symlinked_files_and_directories(tmp_path):
    root, _config, _vault, db, app_db, service = make_index(tmp_path)
    outside = tmp_path / "outside.txt"
    outside.write_text("must not be indexed", encoding="utf-8")
    try:
        (root / "linked.txt").symlink_to(outside)
        (root / "linked-dir").symlink_to(outside.parent, target_is_directory=True)
    except OSError as exc:
        pytest.skip(f"symlinks are unavailable on this system: {exc}")

    result = service.reconcile()

    assert result["discovered"] == 0
    assert db.fetch_one("SELECT id FROM files WHERE filename='linked.txt'") is None
    close_index(db, app_db)
