import fitz
import docx
from contextvault.parsers.plaintext import PlainTextParser
from contextvault.parsers.pdf import PDFParser
from contextvault.parsers.docx import DOCXParser
from contextvault.parsers.spreadsheet import SpreadsheetParser
from contextvault.parsers.image import ImageOCRParser
from contextvault.parsers.registry import create_default_registry

def test_plaintext_parser(tmp_path):
    txt_file = tmp_path / "notes.md"
    txt_file.write_text("# Chapter 1: Operating Systems\n\nCPU scheduling handles process allocation.", encoding="utf-8")
    
    parser = PlainTextParser()
    assert parser.can_parse(".md") is True
    assert parser.can_parse(".txt") is True
    
    doc = parser.parse(txt_file, "file-1")
    assert "Operating Systems" in (doc.title or "")
    assert "CPU scheduling" in doc.text

def test_pdf_parser(tmp_path):
    pdf_path = tmp_path / "sample.pdf"
    doc_fitz = fitz.open()
    page = doc_fitz.new_page()
    page.insert_text((50, 72), "Banker's algorithm avoids deadlock by ensuring safe states.")
    doc_fitz.save(str(pdf_path))
    doc_fitz.close()
    
    parser = PDFParser()
    assert parser.can_parse(".pdf") is True
    doc = parser.parse(pdf_path, "pdf-1")
    assert "Banker's algorithm" in doc.text
    assert len(doc.sections) >= 1
    assert doc.sections[0].page == 1

def test_docx_parser(tmp_path):
    docx_path = tmp_path / "sample.docx"
    doc_docx = docx.Document()
    doc_docx.add_heading("Process Management", level=1)
    doc_docx.add_paragraph("A process is a program in execution.")
    doc_docx.save(str(docx_path))
    
    parser = DOCXParser()
    assert parser.can_parse(".docx") is True
    doc = parser.parse(docx_path, "docx-1")
    assert "Process Management" in doc.text
    assert "program in execution" in doc.text

def test_csv_parser(tmp_path):
    csv_path = tmp_path / "data.csv"
    csv_path.write_text("Topic,Score\nDeadlocks,95\nScheduling,88", encoding="utf-8")
    
    parser = SpreadsheetParser()
    assert parser.can_parse(".csv") is True
    doc = parser.parse(csv_path, "csv-1")
    assert "Deadlocks" in doc.text
    assert "Scheduling" in doc.text

def test_default_registry():
    registry = create_default_registry()
    assert registry.get_parser(".pdf") is not None
    assert registry.get_parser(".docx") is not None
    assert registry.get_parser(".txt") is not None
    assert registry.get_parser(".py") is not None
    assert registry.get_parser(".xlsx") is not None
    assert registry.get_parser(".png") is not None

def test_image_parser_uses_ollama_fallback_when_local_ocr_is_unavailable(tmp_path):
    image_path = tmp_path / "scan.png"
    image_path.write_bytes(b"not-a-real-image-but-the-vision-client-is-tested")

    class FakeVisionClient:
        def ocr_image(self, path):
            assert path == image_path
            return "Invoice total: 42"

    parsed = ImageOCRParser(ocr_client=FakeVisionClient()).parse(image_path, "image-1")
    assert parsed.text == "Invoice total: 42"
    assert parsed.metadata["ocr_backend"] == "ollama-vision"
    assert parsed.metadata["ocr_status"] == "text_extracted"
